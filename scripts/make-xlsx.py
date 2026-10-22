#!/usr/bin/env python3
"""Builds sheet-template/IIS-website-content.xlsx from the CSVs written by `npm run export-sheet`.
Upload the .xlsx to Google Drive and open it with Google Sheets (or File > Import in a new sheet)."""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

D = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'sheet-template')
TABS = ['People', 'Projects', 'Papers', 'Media', 'Press', 'Undergrads', 'Areas']
WIDTH = {'name': 26, 'title': 44, 'role': 28, 'bio': 60, 'bioLong': 60, 'blurb': 50, 'description': 60, 'au': 36, 'areas': 30, 'people': 32, 'projects': 32,
         'email': 26, 'website': 30, 'photo': 40, 'orcid': 22, 'joined': 9, 'year': 8, 'note': 36, 'link': 36, 'url': 40, 'video': 36, 'poster': 30, 'thumb': 30, 'file': 30, 'captions': 28, 've': 22, 'now': 30, 'award': 28, 'outlet': 24, 'by': 20, 'date': 14, 'short': 26, 'word': 14, 'color': 12}
WRAP = {'bio', 'bioLong', 'blurb', 'description', 'title', 'au', 'areas', 'people', 'projects'}
LISTS = {('People', 'status'): 'current,alumni', ('People', 'group'): 'PI,Postdoc,PhD,MPhil,Visitor,Other', ('Media', 'kind'): 'poster,video,talk,demo,other'}
BOOLS = {'chia', 'featured', 'contact', 'visible', 'book'}

wb = Workbook()
guide = wb.active; guide.title = 'Start here'
lines = [
    ('Intelligent Interactive Systems: website content', True),
    ('', False),
    ('This sheet feeds the lab website. Edit a row and the site updates itself (usually within an hour).', False),
    ('You reached it from the "Lab members: edit this site" link at the bottom of the website. Please do not share the password or this link outside the group.', False),
    ('', False),
    ('Tabs: People, Projects, Papers, Media, Press, Undergrads, Areas. One row per item. Do not rename the tabs or the column headings.', False),
    ('', False),
    ('PHOTOS: paste a link to your photo in the "photo" column of the People tab.', True),
    ('  - Google Drive: upload the photo, Share > General access > "Anyone with the link", then Copy link and paste it.', False),
    ('  - Or any public https:// image address (jpg, png or webp, under 8 MB). Square crops look best.', False),
    ('', False),
    ('LINKING: people, projects and areas are matched by NAME. Spell them exactly as they appear in their own tab. Separate several with a semicolon ( ; ).', False),
    ('A paper appears on a project page and on its authors\' pages when you put the project title in "projects" and the names in "people".', False),
    ('', False),
    ('Do not change the "id" column: it forms the page address. Leave it empty for new rows.', False),
    ('To hide a row without deleting it, set "visible" to FALSE.', False),
    ('Web addresses must start with https://', False),
    ('If a row has a mistake, the website keeps showing the last good version and the administrator is emailed what to fix. File > Version history can undo any edit.', False),
    ('', False),
    ('ORCID: put your ORCID iD in the "orcid" column of the People tab (it looks like 0000-0002-1825-0097). New papers on your ORCID record are added to the website automatically every week.', True),
    ('To correct an automatic paper, add a row for it in the Papers tab with the same link; your row wins. To remove one, add that row and set "visible" to FALSE.', False),
    ('', False),
    ('UNDERGRADUATE ALUMNI: one row per person in the Undergrads tab: name, year (and an optional note). They appear in a pop-up under People > Alumni.', True),
    ('', False),
    ('RESEARCH AREAS: the "areas" column must use the "short" names listed in the Areas tab. To add an area, add a row there (name, short, color like #0F8B8D) BEFORE using it elsewhere.', True),
]
for i, (t, b) in enumerate(lines, 1):
    c = guide.cell(row=i, column=1, value=t); c.font = Font(bold=b, size=14 if i == 1 else 11)
guide.column_dimensions['A'].width = 130

for tab in TABS:
    rows = list(csv.reader(open(os.path.join(D, tab + '.csv'), encoding='utf-8')))
    ws = wb.create_sheet(tab)
    for r in rows: ws.append(r)
    head = rows[0]
    for ci, h in enumerate(head, 1):
        c = ws.cell(row=1, column=ci); c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='2540FF')
        c.alignment = Alignment(vertical='center')
        ws.column_dimensions[get_column_letter(ci)].width = WIDTH.get(h, 14)
        L = get_column_letter(ci)
        for ri in range(2, max(len(rows), 300) + 1):
            cell = ws.cell(row=ri, column=ci)
            cell.alignment = Alignment(wrap_text=h in WRAP, vertical='top')
            if h == 'date': cell.number_format = '@'
        if (tab, h) in LISTS:
            dv = DataValidation(type='list', formula1='"' + LISTS[(tab, h)] + '"', allow_blank=True); ws.add_data_validation(dv); dv.add(f'{L}2:{L}500')
        elif h in BOOLS:
            dv = DataValidation(type='list', formula1='"TRUE,FALSE"', allow_blank=True); ws.add_data_validation(dv); dv.add(f'{L}2:{L}500')
        if h == 'id':
            for ri in range(2, max(len(rows), 300) + 1): ws.cell(row=ri, column=ci).font = Font(color='999999')
    ws.freeze_panes = 'B2' if head[0] in ('name', 'title') else 'A2'
    ws.row_dimensions[1].height = 22

out = os.path.join(D, 'IIS-website-content.xlsx'); wb.save(out); print('wrote', out)
